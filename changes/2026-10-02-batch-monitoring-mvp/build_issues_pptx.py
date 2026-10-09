"""Build model_monitoring_issues.pptx from issues.md and the team template.

issues.md is the source. The deck is a view of it in the team's task template
(ai_task_template.pptx), two slides with one tracker table each:

- Slide 1, the project timeline: the issues that are not coding tasks (Decision,
  Discovery, Governance, Infrastructure, Verification), grouped by phase.
- Slide 2, the coding tasks: the Feature issues of all sprints, grouped by sprint.

Each table fits one slide: smaller cell margins and a wider Task column than the template,
and a smaller font when a table has many rows.

PIC: the PIC of the issue's workstream on slide 1 of the template, else the PIC in
"Overview by phase". Plan Start and Plan Finish: the dates of the sprint, or the decision
date of a governance gate. Status, PIC, Plan Start, Plan Finish and Blocker that are
changed in the deck are kept when the deck is built again, matched by issue ID.

PowerPoint files are not committed (.gitignore). Put the team template in this folder as
ai_task_template.pptx before you run the script. The deck is written next to it.

Run from the repository root (python-pptx is not a project dependency):

    uv run --no-project --with python-pptx python changes/2026-10-02-batch-monitoring-mvp/build_issues_pptx.py
"""
from __future__ import annotations

import copy
import re
import sys
from datetime import date
from pathlib import Path

from lxml import etree
from pptx import Presentation

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "issues.md"
TEMPLATE = HERE / "ai_task_template.pptx"
TARGET = HERE / "model_monitoring_issues.pptx"

PHASES = ["Define", "Prepare", "Build", "Validate", "Conclude"]
SPRINT_DATES = {  # Monday to Friday of each sprint
    "Sprint 1": (date(2026, 10, 5), date(2026, 10, 9)),
    "Sprint 2": (date(2026, 10, 12), date(2026, 10, 16)),
    "Sprint 3": (date(2026, 10, 19), date(2026, 10, 23)),
    "Sprint 4": (date(2026, 10, 26), date(2026, 10, 30)),
}
GATES_SECTION = "## Changes on 2026-10-09: AI governance document track"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
SHADES = ("85000", "75000")  # lt1 brightness of the alternating groups, as in the template
TITLE = ("Project Timeline – ", "Model Monitoring App")
CODING_TITLE = ("Coding Tasks – ", "Model Monitoring App")
COLUMNS = ["No", "Phase", "Task", "Status", "PIC", "Plan Start", "Plan Finish", "Blocker"]
KEPT = ["Status", "PIC", "Plan Start", "Plan Finish", "Blocker"]
ID = r"S\d-\d{2}[a-z]?"
ISSUE_ID = re.compile(rf"\b{ID}\b")
ISSUE_HEADING = re.compile(rf"^### ({ID}) — (.+)$")
SPRINT_HEADING = re.compile(r"^## (Sprint \d)\b")

# One-slide fit, between the title and the logo: (most rows, font in hundredths of a
# point, top and bottom cell margin in EMU, table top in EMU). The template uses 800,
# 45720 and 985320.
FITS = [(26, "800", "27432", 985320), (42, "700", "9144", 800000)]
MARGIN_LR = "54864"  # EMU; the template uses 91440
WIDTHS = [380000, 760000, 4900000, 700000, 1500000, 800000, 800000, 2131800]  # sum 11971800


def fmt(d: date) -> str:
    return d.strftime("%d %b %Y")


# --- Parse issues.md ---------------------------------------------------------------------

def plain(text: str) -> str:
    """Strip inline Markdown: links, bold, strikethrough, code spans."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("~~", "").replace("`", "").strip()


def table_cells(line: str) -> list[str]:
    return [plain(c) for c in line.strip().strip("|").split("|")]


def read_table(lines: list[str], start: int) -> list[list[str]]:
    """The body rows of the first Markdown table at or after `start`."""
    i = start
    while not lines[i].lstrip().startswith("|"):
        i += 1
    rows = []
    i += 2  # skip the header and the separator
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        rows.append(table_cells(lines[i]))
        i += 1
    return rows


def heading_at(lines: list[str], heading: str) -> int:
    for i, line in enumerate(lines):
        if line.startswith(heading):
            return i
    raise ValueError(f"issues.md has no section {heading!r}")


def parse(md: str) -> tuple[list[dict], list[dict], list[list[str]]]:
    """The issues, the workstreams of 'Overview by phase', and the governance gates."""
    lines = md.splitlines()
    issues, sprint = [], None
    for k, line in enumerate(lines):
        m = SPRINT_HEADING.match(line)
        if m:
            sprint = m.group(1)
        m = ISSUE_HEADING.match(line)
        if m and sprint:
            meta = read_table(lines, k + 1)[0]  # Type | Owner | Depends on | Tested with | Size
            issues.append({"ID": m.group(1), "Title": plain(m.group(2)), "Type": meta[0],
                           "Sprint": sprint})
    workstreams = [{"Phase": p, "Workstream": w, "IDs": ISSUE_ID.findall(ids), "PIC": pic}
                   for p, w, ids, pic in read_table(lines, heading_at(lines, "## Overview by phase"))]
    gates = read_table(lines, heading_at(lines, GATES_SECTION))
    return issues, workstreams, gates


def workstream_name(task: str) -> str:
    """The workstream name without the issue list at the end of a template Task cell."""
    return re.sub(rf" \((({ID})(, )?)+\)$", "", task)


# --- Read the existing deck and the template ---------------------------------------------

def cell_text(tc) -> str:
    return "".join(t.text or "" for t in tc.iter(A + "t")).strip()


def table_rows(path: Path, first_slide_only: bool = False):
    """Each data row of each table in the deck, as {column: text}."""
    slides = list(Presentation(path).slides)
    for slide in slides[:1] if first_slide_only else slides:
        for shape in slide.shapes:
            if shape.has_table:
                for tr in list(shape._element.iter(A + "tr"))[1:]:
                    yield dict(zip(COLUMNS, (cell_text(tc) for tc in tr.findall(A + "tc"))))


def template_pics() -> dict[str, str]:
    """PIC per workstream name, from slide 1 of the template."""
    return {workstream_name(r["Task"]): r["PIC"]
            for r in table_rows(TEMPLATE, first_slide_only=True) if r.get("Task") and r.get("PIC")}


def kept_values() -> dict[str, dict[str, str]]:
    """Status, PIC, dates and Blocker per issue ID, from the existing deck."""
    if not TARGET.exists():
        return {}
    kept = {}
    for values in table_rows(TARGET):
        m = re.match(rf"^({ID}) ", values.get("Task", ""))
        if m:
            kept[m.group(1)] = {c: values.get(c, "") for c in KEPT}
    return kept


# --- Rows --------------------------------------------------------------------------------

def due_date(text: str) -> date:
    """'Tue 13 Oct' -> date(2026, 10, 13)."""
    day, month = re.search(r"(\d{1,2}) (\w{3})", text).groups()
    return date(2026, ["Oct", "Nov"].index(month) + 10, int(day))


def is_coding(issue: dict) -> bool:
    """A coding task is an issue of type Feature (also "Decision + Feature")."""
    return "Feature" in issue["Type"]


def issue_key(issue: dict) -> tuple:
    sprint, rest = issue["ID"][1:].split("-")
    return int(sprint), rest


def build_slides(md: str) -> list[tuple[tuple[str, str], str, list]]:
    """(title, first-column header, groups) for each slide; a group is (label, rows)."""
    issues, workstreams, gates = parse(md)
    pics, kept = template_pics(), kept_values()
    phase_of, pic_of = {}, {}
    for w in workstreams:
        for i in w["IDs"]:
            phase_of[i] = w["Phase"]
            pic_of[i] = pics.get(w["Workstream"], w["PIC"])
    missing = [i["ID"] for i in issues if i["ID"] not in phase_of]
    if missing:
        raise ValueError(f"not in 'Overview by phase' of issues.md: {', '.join(missing)}")

    finish_of = {}
    for gate in gates:
        for i in ISSUE_ID.findall(gate[4]):
            finish_of.setdefault(i, due_date(gate[5]))

    def row(i: dict) -> dict:
        start, finish = SPRINT_DATES[i["Sprint"]]
        values = {"Task": f'{i["ID"]} {i["Title"]}', "Status": "", "PIC": pic_of[i["ID"]],
                  "Plan Start": fmt(start),
                  "Plan Finish": fmt(finish_of.get(i["ID"], finish)), "Blocker": ""}
        values.update({k: v for k, v in kept.get(i["ID"], {}).items() if v})
        return values

    issues = sorted(issues, key=issue_key)
    by_phase = [(phase, [row(i) for i in issues
                         if not is_coding(i) and phase_of[i["ID"]] == phase])
                for phase in PHASES]
    by_sprint = []
    for sprint, (start, finish) in SPRINT_DATES.items():
        rows = [row(i) for i in issues if is_coding(i) and i["Sprint"] == sprint]
        by_sprint.append((f"{sprint}\n{start.day}–{finish.day} {start:%b}", rows))
    return [
        (TITLE, "Phase", [g for g in by_phase if g[1]]),
        (CODING_TITLE, "Sprint", [g for g in by_sprint if g[1]]),
    ]


# --- Write the tables --------------------------------------------------------------------

def set_text(tc, text: str, run_proto) -> None:
    p = tc.find(f"{A}txBody/{A}p")
    for r in p.findall(A + "r") + p.findall(A + "br"):
        p.remove(r)
    end = p.find(A + "endParaRPr")
    for k, part in enumerate(text.split("\n") if text else []):
        if k:
            br = etree.Element(A + "br")
            br.append(copy.deepcopy(run_proto.find(A + "rPr")))
            end.addprevious(br)
        run = copy.deepcopy(run_proto)
        run.find(A + "t").text = part
        end.addprevious(run)


def compact(tbl, font_size: str, margin_tb: str) -> None:
    """Font size, cell margins and column widths that make the table fit one slide."""
    for el in tbl.iter(A + "rPr", A + "endParaRPr"):
        el.set("sz", font_size)
    for pr in tbl.iter(A + "tcPr"):
        pr.set("marT", margin_tb)
        pr.set("marB", margin_tb)
        pr.set("marL", MARGIN_LR)
        pr.set("marR", MARGIN_LR)
    for col, w in zip(tbl.find(A + "tblGrid").findall(A + "gridCol"), WIDTHS):
        col.set("w", str(w))


def fill_table(tbl, first_header: str, groups: list) -> None:
    rows = tbl.findall(A + "tr")
    header, proto = rows[0], copy.deepcopy(rows[1])
    run_proto = copy.deepcopy(proto.find(f".//{A}r"))
    set_text(header.findall(A + "tc")[1], first_header,
             copy.deepcopy(header.find(f".//{A}r")))
    for tr in rows[1:]:
        tbl.remove(tr)

    n = 0
    for g, (label, items) in enumerate(groups):
        for k, values in enumerate(items):
            n += 1
            tr = copy.deepcopy(proto)
            cells = tr.findall(A + "tc")
            texts = {"No": str(n), "Phase": label if k == 0 else "", **values}
            for col, tc in zip(COLUMNS, cells):
                set_text(tc, texts.get(col, ""), run_proto)
                fill = tc.find(f"{A}tcPr/{A}solidFill")
                if fill is not None:
                    for child in list(fill):
                        fill.remove(child)
                    clr = etree.SubElement(fill, A + "schemeClr", val="lt1")
                    etree.SubElement(clr, A + "lumMod", val=SHADES[g % 2])
            merged = cells[1]
            for attr in ("rowSpan", "vMerge"):
                merged.attrib.pop(attr, None)
            if k == 0 and len(items) > 1:
                merged.set("rowSpan", str(len(items)))
            elif k > 0:
                merged.set("vMerge", "1")
            tbl.append(tr)


def build(slides: list) -> Presentation:
    pres = Presentation(TEMPLATE)
    base = pres.slides[0]
    title_el = copy.deepcopy(base.shapes[0]._element)
    table_el = copy.deepcopy(next(s for s in base.shapes if s.has_table)._element)

    # Keep only the first template slide; the second is another project's tracker.
    id_list = pres.slides._sldIdLst
    for sld in list(id_list)[1:]:
        pres.part.drop_rel(sld.rId)
        id_list.remove(sld)

    for index, (title, first_header, groups) in enumerate(slides):
        count = sum(len(rows) for _, rows in groups)
        fit = next((f for f in FITS if count <= f[0]), None)
        if fit is None:
            raise ValueError(f"{count} rows do not fit one slide (at most {FITS[-1][0]})")
        if index == 0:
            slide = base
        else:
            slide = pres.slides.add_slide(base.slide_layout)
            tree = slide.shapes._spTree
            for shape in list(slide.shapes):
                tree.remove(shape._element)
            tree.append(copy.deepcopy(title_el))
            tree.append(copy.deepcopy(table_el))
        runs = slide.shapes[0].text_frame.paragraphs[0].runs
        runs[0].text, runs[1].text = title
        table = next(s for s in slide.shapes if s.has_table)
        table.top = fit[3]
        tbl = table._element.find(f".//{A}tbl")
        fill_table(tbl, first_header, groups)
        compact(tbl, fit[1], fit[2])
    return pres


def main() -> int:
    if not TEMPLATE.exists():
        print(f"Put the team template at {TEMPLATE}", file=sys.stderr)
        return 1
    slides = build_slides(SOURCE.read_text(encoding="utf-8"))
    build(slides).save(TARGET)
    counts = [sum(len(rows) for _, rows in groups) for _, _, groups in slides]
    print(f"{TARGET.name}: {len(slides)} slides, issues per slide {counts}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
