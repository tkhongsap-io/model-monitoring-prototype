"""Build model_monitoring_issues.xlsx from issues.md.

issues.md is the source. The workbook is a view of it for tracking. Status and Notes (the
yellow cells) are kept from the existing workbook, matched by issue ID.

Run from the repository root (openpyxl is not a project dependency):

    uv run --no-project --with openpyxl python changes/2026-10-02-batch-monitoring-mvp/build_issues_xlsx.py
"""
from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "issues.md"
TARGET = HERE / "model_monitoring_issues.xlsx"

STATUSES = ["To do", "In progress", "Blocked", "Done"]
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
ID_FILL = PatternFill("solid", fgColor="DEEAF6")
EDIT_FILL = PatternFill("solid", fgColor="FFFF99")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
BODY = Font(size=10)
BOLD = Font(size=10, bold=True)
HEAD = Font(size=10, bold=True, color="FFFFFF")

# Tracker sheet: the team's phase template, one row for each workstream in the
# "Overview by phase" table of issues.md.
TRACKER = "Tracker"
TRACKER_STATUSES = ["Plan", "In progress", "Blocked", "Done"]
PHASES = ["Define", "Prepare", "Build", "Validate", "Conclude"]
SPRINT_DATES = {  # Monday to Friday of each sprint
    "Sprint 1": (date(2026, 10, 5), date(2026, 10, 9)),
    "Sprint 2": (date(2026, 10, 12), date(2026, 10, 16)),
    "Sprint 3": (date(2026, 10, 19), date(2026, 10, 23)),
    "Sprint 4": (date(2026, 10, 26), date(2026, 10, 30)),
}
NAVY = PatternFill("solid", fgColor="0E1A47")
SHADES = [PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="BFBFBF")]
WHITE = Side(style="medium", color="FFFFFF")
WHITE_BORDER = Border(left=WHITE, right=WHITE, top=WHITE, bottom=WHITE)
INK = Font(name="Arial", size=9, color="0E1A47")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

ISSUE_HEADING = re.compile(r"^### (S\d-\d{2}) — (.+)$")
SPRINT_HEADING = re.compile(r"^## (Sprint \d)\b")


# --- Markdown to plain text --------------------------------------------------------------

def plain(text: str) -> str:
    """Strip inline Markdown: links, bold, code spans."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?(\s*:?-+:?\s*\|)+\s*:?-*:?\s*\|?\s*", line))


def table_cells(line: str) -> list[str]:
    return [plain(c.strip()) for c in line.strip().strip("|").split("|")]


def block_text(lines: list[str]) -> str:
    """Flatten a Markdown block for one cell: tables become 'a | b', checkboxes become '•'."""
    out: list[str] = []
    for line in lines:
        if line.strip().startswith("```") or is_separator(line):
            continue
        if line.lstrip().startswith("|"):
            out.append(" | ".join(table_cells(line)))
            continue
        line = re.sub(r"^(\s*)- \[[ x]\] ", r"\1• ", line)
        out.append(plain(line))
    while out and not out[0].strip():
        out.pop(0)
    while out and not out[-1].strip():
        out.pop()
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out))


def read_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]]]:
    """Read the Markdown table that starts at or after `start`."""
    i = start
    while not lines[i].lstrip().startswith("|"):
        i += 1
    header = table_cells(lines[i])
    rows = []
    i += 2  # skip the separator
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        rows.append(table_cells(lines[i]))
        i += 1
    return header, rows


def find(lines: list[str], predicate, start: int = 0) -> int:
    for i in range(start, len(lines)):
        if predicate(lines[i]):
            return i
    raise ValueError("expected section not found in issues.md")


# --- Parse issues.md ---------------------------------------------------------------------

def parse(md: str) -> dict:
    lines = md.splitlines()

    _, versions = read_table(lines, find(lines, lambda l: l.startswith("## Fixed versions")))
    _, long_lead = read_table(lines, find(lines, lambda l: l.startswith("**Long-lead requests")))
    _, overview = read_table(lines, find(lines, lambda l: l.startswith("## Overview by phase")))

    issues, paired = [], []
    sprint = None
    i = 0
    while i < len(lines):
        line = lines[i]
        m = SPRINT_HEADING.match(line)
        if m:
            sprint = m.group(1)
        if sprint and line.startswith("**Paired tests this sprint**"):
            _, rows = read_table(lines, i)
            paired += [[sprint, *r] for r in rows]
        m = ISSUE_HEADING.match(line)
        if m:
            end = i + 1
            while end < len(lines) and not (lines[end].startswith("### ")
                                            or lines[end].startswith("## ")
                                            or lines[end].strip() == "---"):
                end += 1
            issues.append(parse_issue(sprint, m.group(1), plain(m.group(2)), lines[i + 1:end]))
            i = end
            continue
        i += 1
    workstreams = [{"Phase": p, "Workstream": w, "IDs": [x.strip() for x in ids.split(",")],
                    "PIC": pic} for p, w, ids, pic in overview]
    return {"issues": issues, "paired": paired, "long_lead": long_lead, "versions": versions,
            "workstreams": workstreams}


def check_overview(issues: list[dict], workstreams: list[dict]) -> list[str]:
    """Each issue must be in exactly one workstream, and each phase must be known."""
    errors = []
    listed = [i for w in workstreams for i in w["IDs"]]
    known = {i["ID"] for i in issues}
    errors += [f"{i} is not in the overview" for i in sorted(known - set(listed))]
    errors += [f"{i} is in the overview but is not an issue" for i in sorted(set(listed) - known)]
    errors += [f"{i} is in more than one workstream" for i in sorted(set(listed))
               if listed.count(i) > 1]
    errors += [f"unknown phase {w['Phase']!r}" for w in workstreams if w["Phase"] not in PHASES]
    return errors


def parse_issue(sprint: str, issue_id: str, title: str, body: list[str]) -> dict:
    head_at = next(k for k, l in enumerate(body) if l.lstrip().startswith("|"))
    header, rows = read_table(body, head_at)
    meta = dict(zip(header, rows[0]))
    rest = body[head_at + 2 + len(rows):]

    def index_of(marker: str) -> int:
        return next((k for k, l in enumerate(rest) if l.strip() == marker), len(rest))

    ac, test = index_of("**Acceptance criteria**"), index_of("**How to test**")
    return {
        "Sprint": sprint, "ID": issue_id, "Title": title,
        "Type": meta["Type"], "Owner": meta["Owner"], "Depends on": meta["Depends on"],
        "Tested with": meta["Tested with"], "Size": meta["Size"],
        "What to do": block_text(rest[:min(ac, test)]),
        "Acceptance criteria": block_text(rest[ac + 1:test]) if ac < len(rest) else "",
        "How to test": block_text(rest[test + 1:]) if test < len(rest) else "",
    }


# --- Write the workbook ------------------------------------------------------------------

def existing_tracking() -> dict[str, tuple[str, str]]:
    """Status and Notes per issue ID from the current workbook, if it exists."""
    if not TARGET.exists():
        return {}
    ws = load_workbook(TARGET)["Issues"]
    cols = {c.value: c.column for c in ws[1]}
    keep = {}
    for r in range(2, ws.max_row + 1):
        issue_id = ws.cell(r, cols["ID"]).value
        if issue_id:
            keep[issue_id] = (ws.cell(r, cols["Status"]).value or "To do",
                              ws.cell(r, cols["Notes"]).value)
    return keep


def task_text(workstream: dict) -> str:
    return f'{workstream["Workstream"]} ({", ".join(workstream["IDs"])})'


def workstream_name(task: str) -> str:
    """The workstream name without the issue list at the end of the Task cell."""
    return re.sub(r" \((S\d-\d{2}(, )?)+\)$", "", task)


def existing_tracker() -> dict[str, tuple[str, str]]:
    """Status and Blocker per workstream from the current Tracker sheet, if it exists."""
    if not TARGET.exists():
        return {}
    wb = load_workbook(TARGET)
    if TRACKER not in wb.sheetnames:
        return {}
    keep = {}
    for row in wb[TRACKER].iter_rows(min_row=2, values_only=True):
        if row[2]:
            keep[workstream_name(str(row[2]))] = (row[3] or "Plan", row[7])
    return keep


def tracker_sheet(wb, issues: list[dict], workstreams: list[dict]) -> None:
    ws = wb.create_sheet(TRACKER, 1)
    ws.sheet_view.showGridLines = False
    header = ["No", "Phase", "Task", "Status", "PIC", "Plan Start", "Plan Finish", "Blocker"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(1, c)
        cell.font, cell.fill = Font(name="Arial", size=9, bold=True, color="FFFFFF"), NAVY
        cell.alignment, cell.border = CENTER, WHITE_BORDER

    kept = existing_tracker()
    sprint_of = {i["ID"]: i["Sprint"] for i in issues}
    row = 2
    for p, phase in enumerate(PHASES):
        group = [w for w in workstreams if w["Phase"] == phase]
        if not group:
            continue
        first = row
        for ws_item in group:
            start = min(SPRINT_DATES[sprint_of[i]][0] for i in ws_item["IDs"])
            finish = max(SPRINT_DATES[sprint_of[i]][1] for i in ws_item["IDs"])
            status, blocker = kept.get(ws_item["Workstream"], ("Plan", None))
            values = [row - 1, phase, task_text(ws_item), status, ws_item["PIC"],
                      start, finish, blocker]
            for c, value in enumerate(values, start=1):
                cell = ws.cell(row, c, value)
                cell.font, cell.fill, cell.border = INK, SHADES[p % 2], WHITE_BORDER
                cell.alignment = LEFT if c in (3, 5, 8) else CENTER
                if c in (6, 7):
                    cell.number_format = "d mmm yyyy"
            row += 1
        ws.merge_cells(start_row=first, start_column=2, end_row=row - 1, end_column=2)

    dv = DataValidation(type="list", formula1='"' + ",".join(TRACKER_STATUSES) + '"',
                        allow_blank=False)
    dv.add(f"D2:D{row - 1}")
    ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGH", [6, 14, 62, 13, 30, 13, 13, 45]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"


def style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEADER_FILL, WRAP, BORDER


def simple_sheet(wb, name: str, header: list[str], rows: list[list[str]], widths: list[int]):
    ws = wb.create_sheet(name)
    ws.append(header)
    for row in rows:
        ws.append(row)
    style_header(ws, len(header))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
    for k, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + k)].width = w
    ws.freeze_panes = "A2"
    return ws


def build(data: dict) -> Workbook:
    issues = data["issues"]
    n = len(issues)
    last = n + 1
    wb = Workbook()

    readme = wb.active
    readme.title = "Read me"
    for text, font in [
        ("Batch Monitoring MVP — Issues", Font(size=14, bold=True)),
        ("Source: changes/2026-10-02-batch-monitoring-mvp/issues.md (reviewed 2026-10-03 and "
         "2026-10-04, revised 2026-10-05: SSO login and the LiteLLM judge). Language: "
         "ASD-STE100, about 80% strict.", BODY),
        (None, BODY),
        ("Sheets", BOLD),
        (f"Tracker — the overall view: {len(data['workstreams'])} workstreams in the team's "
         "phases (Define, Prepare, Build, Validate, Conclude), each with its issue IDs, PIC "
         "and the dates of its sprints. The groups come from 'Overview by phase' in "
         "issues.md. Change only Status and Blocker; the script keeps them. This Status is "
         "separate from the Status of each issue on the Issues sheet.", BODY),
        (f"Issues — one row for each of the {n} issues: owner, dependencies, test partner, "
         "size, what to do, acceptance criteria, how to test.", BODY),
        ("Summary — the number of issues for each sprint, size and status. The numbers are "
         "formulas, so they change when you change the Status column.", BODY),
        ("Paired tests — for each sprint, the feature, its test partner and what the test "
         "proves.", BODY),
        ("Long-lead requests — requests to other teams that must start in Sprint 1.", BODY),
        ("Fixed versions — the exact versions of the components.", BODY),
        (None, BODY),
        ("How to use", BOLD),
        ("Change only the yellow cells: Status (select from the list) and Notes on the Issues "
         "sheet. Do not change the other cells; change issues.md and run "
         "build_issues_xlsx.py. The script keeps the Status and the Notes.", BODY),
        ('Status values: To do, In progress, Blocked, Done. Example: S1-01 → In progress, '
         'with a note such as "Draft schema sent to backend on 6 Oct".', BODY),
        ("Sizes: S = less than 1 day, M = 1 to 3 days, L = 3 to 5 days.", BODY),
    ]:
        readme.append([text])
        readme.cell(readme.max_row, 1).font = font
        readme.cell(readme.max_row, 1).alignment = WRAP
    readme.column_dimensions["A"].width = 120

    summary = wb.create_sheet("Summary")
    summary.append(["Summary"])
    summary["A1"].font = Font(size=14, bold=True)
    summary.append(["All numbers are COUNTIFS formulas on the Issues sheet."])
    summary["A2"].font = BODY
    summary.append([])
    header = ["Sprint", "Issues", "Size S", "Size M", "Size L", "Other size",
              "To do", "In progress", "Blocked", "Done", "Done %"]
    summary.append(header)
    rng = lambda col: f"Issues!${col}$2:${col}${last}"  # noqa: E731
    for k in range(4):
        r = 5 + k
        summary.append([
            f"Sprint {k + 1}",
            f"=COUNTIFS({rng('A')},$A{r})",
            *[f'=COUNTIFS({rng("A")},$A{r},{rng("H")},"{s}")' for s in ("S", "M", "L")],
            f"=B{r}-C{r}-D{r}-E{r}",
            *[f'=COUNTIFS({rng("A")},$A{r},{rng("L")},"{s}")' for s in STATUSES],
            f"=IF(B{r}=0,0,J{r}/B{r})",
        ])
    summary.append(["Total", *[f"=SUM({c}5:{c}8)" for c in "BCDEFGHIJ"], "=IF(B9=0,0,J9/B9)"])
    summary.append(['"Other size" counts issues with a size that is not S, M or L '
                    '(S1-05 is "M + M").'])
    summary["A10"].font = BODY
    for c in range(1, 12):
        cell = summary.cell(4, c)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEADER_FILL, WRAP, BORDER
        for r in range(5, 10):
            cell = summary.cell(r, c)
            cell.font = BOLD if c == 1 else BODY
            cell.border = BORDER
            if c == 11:
                cell.number_format = "0%"
    for col, w in zip("ABCDEFGHIJK", [16, 9, 9, 9, 9, 11, 9, 12, 10, 9, 9]):
        summary.column_dimensions[col].width = w

    columns = ["Sprint", "ID", "Title", "Type", "Owner", "Depends on", "Tested with", "Size",
               "What to do", "Acceptance criteria", "How to test", "Status", "Notes"]
    tracking = existing_tracking()
    ws = wb.create_sheet("Issues")
    ws.append(columns)
    for issue in issues:
        status, notes = tracking.get(issue["ID"], ("To do", None))
        ws.append([issue[c] for c in columns[:11]] + [status, notes])
    style_header(ws, len(columns))
    for r in range(2, last + 1):
        for c in range(1, len(columns) + 1):
            cell = ws.cell(r, c)
            cell.font = BOLD if c == 2 else BODY
            cell.alignment, cell.border = WRAP, BORDER
            if c <= 2:
                cell.fill = ID_FILL
            elif c >= 12:
                cell.fill = EDIT_FILL
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    dv.add(f"L2:L{last}")
    ws.add_data_validation(dv)
    for col, w in zip("ABCDEFGHIJKLM", [10, 8, 34, 14, 22, 22, 22, 7, 70, 60, 60, 13, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{last}"

    simple_sheet(wb, "Paired tests", ["Sprint", "Feature", "Tested with", "The test proves",
                                      "Environment"], data["paired"], [10, 30, 30, 60, 24])
    simple_sheet(wb, "Long-lead requests", ["Request", "Owner", "Needed by",
                                            "If it is not approved in time"],
                 data["long_lead"], [50, 26, 22, 60])
    simple_sheet(wb, "Fixed versions", ["Component", "Image or package", "Version", "Note"],
                 data["versions"], [28, 40, 22, 70])
    tracker_sheet(wb, issues, data["workstreams"])
    return wb


def main() -> int:
    data = parse(SOURCE.read_text(encoding="utf-8"))
    ids = [i["ID"] for i in data["issues"]]
    if len(ids) != len(set(ids)):
        print("duplicate issue IDs in issues.md", file=sys.stderr)
        return 1
    errors = check_overview(data["issues"], data["workstreams"])
    if errors:
        print("issues.md, Overview by phase:\n  " + "\n  ".join(errors), file=sys.stderr)
        return 1
    build(data).save(TARGET)
    print(f"{TARGET.name}: {len(ids)} issues, {len(data['paired'])} paired tests, "
          f"{len(data['long_lead'])} long-lead requests, {len(data['versions'])} versions")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
